from __future__ import annotations

import ast
import hashlib
import json
import re
import shutil
import time
from dataclasses import asdict, dataclass
from pathlib import Path

from .core import Budget, Events, digest
from .items import CARD_CHARS, Item

SUPPORT_VERSION = 2
OBSERVATION_RESULT_PREFIX = "__ROUGE_OBSERVATION_RESULT_V2__="
PRODUCER_MARKER = "ROUGE_INPUT_PRODUCER_V1"
from .demands import FACT_KINDS, FACT_USE


SUPPORT_CATALOG_VERSION = 3
SUPPORT_KINDS = ("structure", "typed_values") + FACT_KINDS
PRODUCER_SCHEMA = (
    "Define exactly def observe(books): returning JSON-compatible facts. books is a list of "
    "input-only dictionaries: {file_index, sheets: [{name, rows, columns, cells: "
    "[{row, column, coordinate, value, kind, number_format}]}]}. value is null, bool, number "
    "or a string of at most 120 characters; kind is the Excel data_type. rows and columns "
    "are integer dimensions, not row/cell lists; iterate sheet['cells'] for cell records. The driver samples "
    "at most 3 cases, 4 sheets, first 64 rows and 20 columns. Use only pure Python: no imports, "
    "file APIs, private attributes, external state or writes to workbooks. Return observations "
    "or normalized input facts, never answers. Your source plus its inspect-card JSON wrapper "
    f"must fit {CARD_CHARS} characters."
)


class ObservationFailure(ValueError):
    pass

_SOURCES = {
    "structure": '''def observe(books):
    out = []
    for book in books:
        sheets = []
        for sheet in book["sheets"]:
            headers = [c for c in sheet["cells"] if c["row"] <= 6 and c["value"] is not None]
            sheets.append({"name": sheet["name"], "rows": sheet["rows"], "columns": sheet["columns"], "header_rows": headers[:24]})
        out.append({"file_index": book["file_index"], "sheet_map": sheets})
    return out
''',
    "typed_values": '''def observe(books):
    out = []
    for book in books:
        sheets = []
        for sheet in book["sheets"]:
            cells = [c for c in sheet["cells"] if c["value"] is not None]
            special = [c for c in cells if c["kind"] != "s" or (isinstance(c["value"], str) and ("(" in c["value"] or "," in c["value"]))]
            sheets.append({"name": sheet["name"], "typed_values": special[:24], "sample_values": cells[:8]})
        out.append({"file_index": book["file_index"], "sheets": sheets})
    return out
''',
}

_SOURCES.update({
    'column_profile': '''def observe(books):
    out = []
    for book in books:
        sheets = []
        for sheet in book["sheets"][:2]:
            cols = []
            for col in range(1, min(sheet["columns"], 8) + 1):
                cells = [c for c in sheet["cells"] if c["column"] == col and c["value"] is not None]
                headers = [{"row": c["row"], "text": str(c["value"])[:40]} for c in cells if c["row"] <= 6][:3]
                kinds = {}
                for c in cells:
                    key = c["kind"]
                    kinds[key] = kinds.get(key, 0) + 1
                cols.append({"column": col, "headers": headers, "kinds": kinds})
            sheets.append({"name": sheet["name"], "column_profile": cols})
        out.append({"file_index": book["file_index"], "sheets": sheets})
    return out
''',
    'numeric_profile': '''def observe(books):
    out = []
    for book in books:
        sheets = []
        for sheet in book["sheets"][:2]:
            values = []
            for c in sheet["cells"]:
                raw = c["value"]
                if raw is None or isinstance(raw, bool) or c["kind"] == "f":
                    continue
                v = str(raw).strip()
                neg = v.startswith("(") and v.endswith(")")
                if neg:
                    v = v[1:-1].strip()
                pct = v.endswith("%")
                if pct:
                    v = v[:-1].strip()
                digits = v[1:] if v[:1] in ("+", "-") else v
                if neg and digits != v:
                    continue
                groups = digits.split(".")[0].split(",")
                if len(groups) > 1 and (not groups[0].isdigit() or not 1 <= len(groups[0]) <= 3 or not all(len(g) == 3 and g.isdigit() for g in groups[1:])):
                    continue
                v = v.replace(",", "")
                digits = v[1:] if v[:1] in ("+", "-") else v
                if digits.replace(".", "", 1).isdigit():
                    number = float(v) * (-1 if neg else 1) / (100 if pct else 1)
                    values.append({"coordinate": c["coordinate"], "raw": raw, "numeric_candidate": number, "number_format": c["number_format"]})
                if len(values) >= 10:
                    break
            sheets.append({"name": sheet["name"], "numeric_profile": values})
        out.append({"file_index": book["file_index"], "sheets": sheets})
    return out
''',
    'layout_profile': '''def observe(books):
    out = []
    for book in books:
        layouts = []
        for sheet in book["sheets"][:3]:
            headers = [{"row": c["row"], "column": c["column"], "text": str(c["value"])[:40]} for c in sheet["cells"] if c["row"] <= 3 and c["value"] is not None][:8]
            layouts.append({"name": sheet["name"], "rows": sheet["rows"], "columns": sheet["columns"], "header_signature": headers})
        out.append({"file_index": book["file_index"], "layout_profile": layouts})
    return out
''',
})

_BUILTINS = {"len", "str", "int", "float", "bool", "list", "dict", "tuple", "set",
             "min", "max", "sum", "range", "enumerate", "zip", "sorted", "round", "abs",
             "isinstance", "reversed", "any", "all"}
_METHODS = {"append", "extend", "get", "items", "keys", "values", "strip", "lstrip", "rstrip",
            "replace", "split", "rsplit", "lower", "upper", "casefold", "startswith", "endswith",
            "isdigit", "isnumeric", "isdecimal", "join", "count", "index"}
_NODES = (ast.Module, ast.FunctionDef, ast.arguments, ast.arg, ast.Return, ast.Assign, ast.AugAssign,
          ast.For, ast.If, ast.Expr, ast.Pass, ast.Break, ast.Continue, ast.Name, ast.Load, ast.Store,
          ast.Constant, ast.List, ast.Tuple, ast.Dict, ast.Set, ast.Subscript, ast.Slice, ast.Call,
          ast.Attribute, ast.keyword, ast.Compare, ast.BoolOp, ast.UnaryOp, ast.BinOp, ast.IfExp,
          ast.ListComp, ast.SetComp, ast.DictComp, ast.GeneratorExp, ast.comprehension,
          ast.And, ast.Or, ast.Not, ast.UAdd, ast.USub, ast.Add, ast.Sub, ast.Mult, ast.Div,
          ast.FloorDiv, ast.Mod, ast.Eq, ast.NotEq, ast.Lt, ast.LtE, ast.Gt, ast.GtE,
          ast.In, ast.NotIn, ast.Is, ast.IsNot)


def validate_producer_source(source: str) -> None:

    if not isinstance(source, str) or not 1 <= len(source) <= CARD_CHARS:
        raise ValueError(f"Producer source must contain 1..{CARD_CHARS} characters")
    try:
        tree = ast.parse(source)
    except SyntaxError as exc:
        location = f" at line {exc.lineno}, column {exc.offset}"
        excerpt = (exc.text or '').strip()[:200]
        raise ValueError(f"Invalid producer Python: {exc.msg}{location}; source: {excerpt}") from None
    if len(tree.body) != 1 or not isinstance(tree.body[0], ast.FunctionDef):
        raise ValueError("Producer must contain exactly one function: def observe(books)")
    fn = tree.body[0]
    args = fn.args
    if (fn.name != "observe" or fn.decorator_list or fn.returns is not None or
            len(args.args) != 1 or args.args[0].arg != "books" or args.args[0].annotation is not None or
            args.posonlyargs or args.kwonlyargs or args.vararg or args.kwarg or args.defaults):
        raise ValueError("Producer signature must be exactly def observe(books)")
    if not any(isinstance(node, ast.Name) and node.id == "books" and isinstance(node.ctx, ast.Load)
               for node in ast.walk(fn)):
        raise ValueError("Producer must derive its observations from books")
    for node in ast.walk(tree):
        if not isinstance(node, _NODES):
            raise ValueError(f"Producer cannot use {type(node).__name__}")
        if isinstance(node, ast.FunctionDef) and node is not fn:
            raise ValueError("Nested functions are not allowed in a producer")
        if isinstance(node, ast.Name) and node.id.startswith("_"):
            raise ValueError("Private names are not allowed in a producer")
        if isinstance(node, ast.Attribute):
            if node.attr not in _METHODS or isinstance(node.ctx, ast.Store):
                raise ValueError(f"Producer attribute {node.attr!r} is not allowed")
        if isinstance(node, ast.Call):
            if isinstance(node.func, ast.Name):
                if node.func.id not in _BUILTINS:
                    raise ValueError(f"Producer call {node.func.id!r} is not allowed")
            elif not isinstance(node.func, ast.Attribute) or node.func.attr not in _METHODS:
                raise ValueError("Producer may call only approved pure builtins and methods")
        if isinstance(node, ast.Constant) and isinstance(node.value, str):
            value = node.value.lower()
            if (len(value) > 300 or "\\" in value or ".xlsx" in value or
                    any(x in value for x in ("/work", "golden", "recalculated", "case_", "_output")) or
                    value in {"answer", "answers", "expected", "expected_answer", "output", "outputs",
                              "solution", "solutions"} or
                    re.fullmatch(r"\d+-\d+", value) or
                    re.fullmatch(r"\$?[a-z]{1,3}\$?\d+(?::\$?[a-z]{1,3}\$?\d+)?", value) or
                    re.fullmatch(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", value)):
                raise ValueError("Producer may not embed paths, task IDs, cell addresses or answer/output names")


@dataclass(frozen=True)
class Producer:
    source: str
    instructions: str
    name: str

    @property
    def id(self) -> str:
        return digest([SUPPORT_VERSION, self.source, self.instructions])


def format_producer_item(name: str, source: str, instructions: str = "", needs=()) -> Item:
    validate_producer_source(source)
    instructions = instructions or "Use the input-only producer observations; inspect further when needed."
    payload = json.dumps({"source": source, "instructions": instructions}, ensure_ascii=False)
    return Item("inspect", name, f"{PRODUCER_MARKER}\n```json\n{payload}\n```", tuple(needs))


def parse_producer(item: Item) -> Producer | None:
    if not item.content.startswith(PRODUCER_MARKER):
        return None
    if item.slot != "inspect":
        raise ValueError("An input producer must occupy the inspect slot")
    match = re.fullmatch(re.escape(PRODUCER_MARKER) + r"\s*```json\s*([\s\S]*?)\s*```\s*", item.content)
    if not match:
        raise ValueError("Invalid input producer JSON contract")
    payload = json.loads(match.group(1))
    if (not isinstance(payload, dict) or set(payload) != {"source", "instructions"} or
            not isinstance(payload["instructions"], str)):
        raise ValueError("Producer contract requires source and instructions strings")
    validate_producer_source(payload["source"])
    return Producer(payload["source"], payload["instructions"], item.name)


def make_source_item(kind: str) -> Item:
    if kind not in SUPPORT_KINDS:
        raise ValueError(f"Unknown support kind: {kind!r}")
    guidance = 'Fact interface: ' + kind + '. ' + FACT_USE[kind] if kind in FACT_KINDS else ''
    return format_producer_item(f"source_{kind}", _SOURCES[kind], guidance)


def retain_inspection(item: Item, build) -> Item:

    producer = parse_producer(item)
    if producer is None:
        raise ValueError('Permanent completion must be an input producer')
    marker = 'Retained inspection guidance:\n'
    if producer.instructions.startswith(marker):
        return item
    previous = build.get(('inspect', ''))
    old = parse_producer(previous) if previous else None
    guidance = old.instructions if old else previous.content if previous else ''
    if guidance.startswith(marker):
        guidance = guidance[len(marker):].split('\nNew producer guidance:\n', 1)[0]
    if not guidance:
        return item


    remaining = CARD_CHARS - len(marker) - len('\nNew producer guidance:\n') - len(producer.instructions)
    if remaining <= 0:
        raise ValueError('Producer guidance leaves no room for retained inspection')
    guidance = guidance[:remaining]
    return format_producer_item(item.name, producer.source,
                                marker + guidance + '\nNew producer guidance:\n' + producer.instructions,
                                item.needs)


@dataclass(frozen=True)
class SupportObservation:
    kind: str
    source_id: str
    facts: str
    success: bool
    returncode: int
    cost: float
    seconds: float
    input_hashes: dict
    error: str = ""
    truncated: bool = False
    boundary: str = "before_first_actor_call"
    version: int = SUPPORT_VERSION

    def to_dict(self) -> dict:
        return asdict(self)


def _driver(source: str, chars: int) -> str:

    return '''import json
from pathlib import Path
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

books = []
for path in sorted(Path(".").glob("case_*_input.xlsx"), key=lambda p: int(p.name.split("_")[1]))[:3]:
    index = int(path.name.split("_")[1])
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    sheets = []
    try:
        for sheet in workbook.worksheets[:4]:
            cells = []
            for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(64, sheet.max_row or 64), min_col=1, max_col=min(20, sheet.max_column or 20)), 1):
                for column_index, cell in enumerate(row, 1):
                    value = cell.value
                    if isinstance(value, ArrayFormula):
                        value = f"{value.text or ''} [array range: {value.ref}]"
                    elif isinstance(value, DataTableFormula):
                        value = "data table: " + json.dumps(dict(value), sort_keys=True)
                    elif value is not None and not isinstance(value, (str, bool, int, float)):
                        value = str(value)
                    if isinstance(value, str):
                        value = value[:120]
                    coordinate = openpyxl.utils.get_column_letter(column_index) + str(row_index)
                    cells.append({"row": row_index, "column": column_index, "coordinate": coordinate, "value": value, "kind": cell.data_type, "number_format": cell.number_format})
            sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "cells": cells})
        books.append({"file_index": index, "sheets": sheets})
    finally:
        workbook.close()

''' + source + '''
value = observe(books)
facts = json.dumps(value, ensure_ascii=False, allow_nan=False, sort_keys=True)
limit = ''' + str(chars) + '''
truncated = len(facts) > limit
if truncated:
    # Keep complete scalar records, never cut a JSON value or invent a summary answer.
    records, fields = [], set()
    def collect(node, path):
        if isinstance(node, dict):
            scalars = {str(k): v for k, v in node.items() if not isinstance(v, (list, dict))}
            fields.update(str(k) for k in node)
            if scalars:
                records.append({"path": path, "facts": scalars})
            for key, child in node.items():
                if isinstance(child, (list, dict)):
                    collect(child, path + "." + str(key))
        elif isinstance(node, list):
            for index, child in enumerate(node):
                collect(child, path + "[" + str(index) + "]")
        else:
            records.append({"path": path, "value": node})
    collect(value, "root")
    summary = {"available_fields": sorted(fields)[:20], "records": [],
               "omitted_records": len(records), "truncated": True}
    def encode(data):
        return json.dumps(data, ensure_ascii=False, allow_nan=False, sort_keys=True)
    while len(encode(summary)) > limit and summary["available_fields"]:
        summary["available_fields"].pop()
    for record in records:
        summary["records"].append(record)
        summary["omitted_records"] -= 1
        if len(encode(summary)) > limit:
            summary["records"].pop()
            summary["omitted_records"] += 1
    facts = encode(summary)
print(''' + repr(OBSERVATION_RESULT_PREFIX) + ''' + json.dumps({"facts": facts, "truncated": truncated}, ensure_ascii=False))
'''


def _decode_driver_output(output: str) -> dict:

    records = [line[len(OBSERVATION_RESULT_PREFIX):] for line in output.splitlines()
               if line.startswith(OBSERVATION_RESULT_PREFIX)]
    if len(records) != 1:
        raise ValueError(f"Expected one framed observation result, received {len(records)}")
    return json.loads(records[0])


def contains_fact_key(value, kind):

    if isinstance(value, dict):
        return (kind in value or kind in value.get('available_fields', []) or
                any(contains_fact_key(child, kind) for child in value.values()))
    if isinstance(value, list):
        return any(contains_fact_key(child, kind) for child in value)
    return False


class SupportProvider:

    kinds = SUPPORT_KINDS

    def __init__(self, runner, budget: Budget, events: Events, config: dict):
        self.runner, self.budget, self.events = runner, budget, events
        self.chars = int(config.get("support_chars", 3000))
        if not 128 <= self.chars <= 6000:
            raise ValueError("support_chars must be within 128..6000")

    def observe(self, kind: str, workspace: Path, *, producer: Producer | None = None,
                temporary: bool = True) -> SupportObservation:
        producer = producer or parse_producer(make_source_item(kind))
        if producer is None:
            raise ValueError("An observation needs a producer")
        validate_producer_source(producer.source)
        purpose = "temporary_support" if temporary else "permanent_producer"
        source_work = workspace.parent / "input_sources" / digest([purpose, producer.id, time.time_ns()])
        source_work.mkdir(parents=True, exist_ok=False)
        source_work.chmod(0o777)
        hashes = {}
        for path in sorted(workspace.glob("case_*_input.xlsx")):
            if not path.is_file() or path.is_symlink():
                raise ValueError("Producer inputs must be regular original input workbooks")
            hashes[path.name] = hashlib.sha256(path.read_bytes()).hexdigest()
            shutil.copyfile(path, source_work / path.name)
        script = source_work / "observation.py"
        script.write_text(_driver(producer.source, self.chars), encoding="utf-8")
        start = time.perf_counter()
        code, output = self.runner.run(source_work, ["python", "/work/observation.py"], self.budget, purpose)
        seconds, facts, error, truncated = time.perf_counter() - start, "", "", False
        if code:
            error = f"Observation program failed ({code}): {output[-600:]}"
        else:
            try:
                data = _decode_driver_output(output)
                if not isinstance(data, dict) or not isinstance(data.get("facts"), str):
                    raise ValueError("Invalid observation result contract")
                facts = data["facts"]
                if len(facts) > self.chars:
                    raise ValueError("Observation exceeded its declared character bound")
                decoded = json.loads(facts)
                declared = re.search(r'Fact interface: ([a-z_]+)\.', producer.instructions)
                if declared and declared.group(1) in FACT_KINDS and not contains_fact_key(decoded, declared.group(1)):
                    raise ValueError('Producer did not return its declared fact interface')
                truncated = bool(data.get("truncated"))
            except (ValueError, TypeError) as exc:
                error = f"Invalid observation output: {str(exc)[:300]}"
        result = SupportObservation(kind, producer.id, facts, not error, code,
                                    self.budget.tool_price, seconds, hashes, error, truncated=truncated)
        event = result.to_dict()
        event["support_kind"] = event.pop("kind")
        self.events.emit("input_observation", temporary=temporary, **event)
        return result

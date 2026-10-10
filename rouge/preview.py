from __future__ import annotations

import json
from pathlib import Path

PREVIEW_VERSION = 2


def cell_text(value):

    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
    if value is None:
        return None
    if isinstance(value, ArrayFormula):
        return f"{value.text or ''} [array range: {value.ref}]"[:100]
    if isinstance(value, DataTableFormula):
        return ("data table: " + json.dumps(dict(value), sort_keys=True))[:100]
    return str(value)[:100]


def workbook_preview(paths: list[Path], max_chars: int) -> str:

    import openpyxl
    if max_chars <= 0:
        return ""
    books = []
    for index, path in enumerate(paths[:3]):
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            sheets = []
            for sheet in book.worksheets[:4]:
                rows = [[cell_text(c.value) for c in row]
                        for row in sheet.iter_rows(min_row=1, max_row=8, min_col=1, max_col=12)]
                sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                               "first_rows_A_to_L": rows})
            books.append({"file": f"case_{index}_input.xlsx", "sheet_names": book.sheetnames,
                          "sampled_sheets": sheets})
        finally:
            book.close()
    text = json.dumps(books, ensure_ascii=False, default=str)
    return text[:max_chars] + ("\n[preview truncated; inspect other cells with Python]" if len(text) > max_chars else "")

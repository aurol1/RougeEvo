from __future__ import annotations

import re

REF = re.compile(r"\$?[A-Za-z]{1,3}\$?\d*(?::\$?[A-Za-z]{1,3}\$?\d*)?")
CELL = re.compile(r"([A-Z]{1,3})(\d*)")


def _clean(ref: str) -> str:
    ref = ref.replace("$", "").upper()
    if ":" not in ref and not any(c.isdigit() for c in ref):
        raise ValueError(f"A single reference needs a row: {ref}")
    return ref


def _sheet_suffix(prefix: str, sheetnames: list[str]) -> str | None:

    text = prefix.strip().strip("'").strip()
    found = [s for s in sheetnames if text.endswith(s) or text.endswith(s.strip())]
    return max(found, key=len, default=None)


def parse_targets(position: str, answer_sheet: str | None, sheetnames: list[str]) -> list[tuple[str, str]]:
    text = str(position).replace("\xa0", " ").strip()
    if not text or not sheetnames:
        raise ValueError("Empty answer position or workbook")
    out = []
    if "!" in text:
        chunks = text.split("!")
        for i in range(1, len(chunks)):
            sheet = _sheet_suffix(chunks[i - 1], sheetnames)
            match = REF.match(chunks[i].strip().lstrip("'"))
            if sheet is None or match is None:
                raise ValueError(f"Cannot resolve {text!r} against sheets {sheetnames}")
            out.append((sheet, _clean(match.group(0))))
        return out
    sheet_text = str(answer_sheet or "").replace("\xa0", " ").strip()
    if sheet_text in sheetnames:
        sheets = [sheet_text]
    else:
        listed = [p.strip().strip("'") for p in sheet_text.split(",") if p.strip()]
        sheets = [s for s in sheetnames if s in listed or s.strip() in listed]
    for part in text.split(","):
        match = REF.fullmatch(part.strip().strip("'"))
        if match is None:
            raise ValueError(f"Cannot parse range {part!r} in {text!r}")

        out.extend((s, _clean(match.group(0))) for s in (sheets or [sheetnames[0]]))
    return out


def display(targets) -> str:
    return ",".join("'{}'!{}".format(sheet.replace("'", "''"), ref) for sheet, ref in targets)


def _last_row(sheet, first_col: int, last_col: int) -> int:

    rows = [r for (r, c), cell in getattr(sheet, "_cells", {}).items()
            if first_col <= c <= last_col and cell.value is not None]
    return max(rows, default=1)


def bounded(ref: str, *sheets) -> str:

    from openpyxl.utils import column_index_from_string

    parts = ref.split(":")
    cells = [CELL.fullmatch(p) for p in parts]
    if any(m is None for m in cells):
        raise ValueError(f"Bad range {ref}")
    if all(m.group(2) for m in cells):
        return ref
    cols = [column_index_from_string(m.group(1)) for m in cells]
    lo, hi = min(cols), max(cols)
    last = max(_last_row(s, lo, hi) for s in sheets)
    start = int(cells[0].group(2) or 1)
    end = int(cells[-1].group(2)) if cells[-1].group(2) else max(last, start)
    return f"{cells[0].group(1)}{start}:{cells[-1].group(1)}{end}"

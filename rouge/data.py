from __future__ import annotations

import contextlib
import hashlib
import importlib.util
import io
import json
import threading
from dataclasses import dataclass
from pathlib import Path
from types import SimpleNamespace

from .positions import bounded, display


@dataclass(frozen=True)
class Case:
    input: Path
    answer: Path


@dataclass(frozen=True)
class Task:
    id: str
    instruction: str
    instruction_type: str
    answer_position: str
    cases: tuple[Case, ...]
    targets: tuple[tuple[str, str], ...] | None = None

    def public(self) -> dict:
        return {"id": self.id, **self.agent_view()}

    def agent_view(self) -> dict:

        return {"instruction": self.instruction, "instruction_type": self.instruction_type,
                "answer_position": self.answer_position,
                "files": [f"case_{i}_input.xlsx" for i in range(len(self.cases))],
                "outputs": [f"case_{i}_output.xlsx" for i in range(len(self.cases))]}


class Dataset:
    def __init__(self, manifest: Path, verify_hashes: bool = True):
        self.path = manifest.resolve()
        self.root = self.path.parent
        raw = self.path.read_bytes()
        self.fingerprint = hashlib.sha256(raw).hexdigest()
        self.metadata = json.loads(raw)
        if self.metadata.get("schema_version") != 1:
            raise ValueError("Unsupported manifest schema")
        self.tasks = {}
        for row in self.metadata["tasks"]:
            key = str(row["id"])
            if key in self.tasks:
                raise ValueError(f"Duplicate task: {key}")
            cases = []
            for entry in row["cases"]:
                paths = [self._path(entry[k]) for k in ("input", "answer")]
                if verify_hashes:
                    for k, path in zip(("input", "answer"), paths):
                        self._verify(path, entry[k + "_sha256"])
                cases.append(Case(*paths))
            if not cases:
                raise ValueError(f"No cases: {key}")
            targets = row.get("targets")
            if targets is not None:
                targets = tuple((str(s), str(r)) for s, r in targets)
                if not targets:
                    raise ValueError(f"Empty targets: {key}")
            position = display(targets) if targets else row["answer_position"]
            self.tasks[key] = Task(key, row["instruction"], row["instruction_type"], position,
                                   tuple(cases), targets)
        self.splits = self.metadata["splits"]
        ids = [x for name in ("train", "validation", "test") for x in self.splits[name]]
        if any(not self.splits[x] for x in ("train", "validation", "test")):
            raise ValueError("All three splits must be nonempty")
        if len(set(ids)) != len(ids) or set(ids) != set(self.tasks):
            raise ValueError("Splits must partition task IDs without overlap or omissions")
        self.scorer_path = self._path(self.metadata["scorer"])
        self._verify(self.scorer_path, self.metadata["scorer_sha256"])

    def _path(self, name: str) -> Path:
        path = (self.root / name).resolve()
        try:
            path.relative_to(self.root)
        except ValueError:
            raise ValueError(f"Invalid dataset path: {name}")
        if not path.is_file():
            raise ValueError(f"Invalid dataset path: {name}")
        return path

    @staticmethod
    def _verify(path: Path, expected: str) -> None:
        if hashlib.sha256(path.read_bytes()).hexdigest() != expected:
            raise ValueError(f"Content hash mismatch: {path}")

    def split(self, name: str, limit: int = 0) -> list[Task]:
        ids = self.splits[name]
        if limit:
            ids = ids[:limit]
        return [self.tasks[x] for x in ids]


def _short(value, width: int = 40) -> str:
    text = repr(value)
    return text if len(text) <= width else text[:width - 3] + "..."


class OfficialScorer:

    _lock = threading.Lock()

    def __init__(self, path: Path):
        spec = importlib.util.spec_from_file_location("spreadsheetbench_official", path)
        if spec is None or spec.loader is None:
            raise ValueError("Cannot load pinned official scorer")
        self.module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.module)
        import openpyxl
        original = openpyxl.load_workbook

        def managed_load(filename, **kwargs):

            with Path(filename).open("rb") as source:
                return original(source, **kwargs)

        self.load = managed_load
        self.module.openpyxl = SimpleNamespace(load_workbook=managed_load)

    def compare(self, task: Task, case: Case, output: Path) -> bool:
        return self.check(task, case, output)[0]

    def check(self, task: Task, case: Case, output: Path, explain: bool = False,
              reveal_expected: bool = True) -> tuple[bool, str]:

        if not output.is_file() or output.is_symlink():
            return False, "no output workbook was produced"
        with self._lock, contextlib.redirect_stdout(io.StringIO()):
            if task.targets is None:

                ok, _ = self.module.compare_workbooks(str(case.answer), str(output),
                                                       task.instruction_type, task.answer_position)
                note = "" if ok or not explain else self._legacy_explain(task, case, output,
                                                                           reveal_expected)
                return bool(ok), note
            gt = self.load(case.answer, data_only=True)
            try:
                proc = self.load(output, data_only=True)
            except Exception as exc:
                return False, f"output workbook unreadable ({type(exc).__name__})"
            verdicts, notes = [], []
            for sheet, ref in task.targets:
                if sheet not in proc.sheetnames:
                    verdicts.append(False)
                    notes.append(f"sheet {sheet!r} missing from output")
                    continue
                area = bounded(ref, gt[sheet], proc[sheet])
                ok, _ = self.module.cell_level_compare(gt, proc, sheet, area)
                verdicts.append(bool(ok))
                if not ok and explain:
                    notes.append(self.explain(gt[sheet], proc[sheet], sheet, area, reveal_expected))
        return all(verdicts), "; ".join(notes)[:900]

    def _legacy_explain(self, task: Task, case: Case, output: Path, reveal: bool) -> str:

        try:
            gt = self.load(case.answer, data_only=True)
            proc = self.load(output, data_only=True)
        except Exception as exc:
            return f"output workbook unreadable ({type(exc).__name__})"
        notes = []
        try:
            for part in task.answer_position.split(","):
                sheet, ref = part.split("!") if "!" in part else (gt.sheetnames[0], part)
                sheet, ref = sheet.strip("'"), ref.strip("'")
                if sheet not in proc.sheetnames:
                    notes.append(f"sheet {sheet!r} missing from output")
                    continue
                notes.append(self.explain(gt[sheet], proc[sheet], sheet, bounded(ref.upper(), gt[sheet],
                                                                                 proc[sheet]), reveal))
        except Exception as exc:
            notes.append(f"values differ (diagnostic unavailable: {type(exc).__name__})")
        return "; ".join(notes)[:900]

    def training_progress(self, task: Task, case: Case, output: Path) -> dict:

        from openpyxl.utils.cell import range_boundaries
        gt, proc = None, None
        try:
            with self._lock:
                gt = self.load(case.answer, data_only=True)
                if output.is_file() and not output.is_symlink():
                    try:
                        proc = self.load(output, data_only=True)
                    except Exception:
                        proc = None
                targets = task.targets
                if targets is None:
                    targets = []
                    for part in task.answer_position.split(','):
                        sheet, ref = part.split('!') if '!' in part else (gt.sheetnames[0], part)
                        targets.append((sheet.strip("'"), ref.strip("'").upper()))
                seen, matched = set(), 0
                for sheet, ref in targets:
                    sheets = [gt[sheet]] + ([proc[sheet]] if proc and sheet in proc.sheetnames else [])
                    c1, r1, c2, r2 = range_boundaries(bounded(ref, *sheets))
                    if (r2 - r1 + 1) * (c2 - c1 + 1) > 200000:
                        return {'available': False, 'reason': 'target_range_too_large'}
                    for row in range(r1, r2 + 1):
                        for col in range(c1, c2 + 1):
                            key = (sheet, row, col)
                            if key in seen:
                                continue
                            seen.add(key)
                            if proc is not None and sheet in proc.sheetnames and self.module.compare_cell_value(
                                    gt[sheet].cell(row, col).value, proc[sheet].cell(row, col).value):
                                matched += 1
                total = len(seen)
                return {'available': total > 0, 'matched': matched, 'total': total,
                        'fraction': matched / total if total else 0.0}
        except Exception as exc:
            return {'available': False, 'reason': type(exc).__name__}
        finally:
            for book in (gt, proc):
                if book is not None:
                    book.close()

    def explain(self, gt, proc, sheet: str, area: str, reveal: bool, limit: int = 3) -> str:
        from openpyxl.utils.cell import range_boundaries
        c1, r1, c2, r2 = range_boundaries(area)
        if (c2 - c1 + 1) * (r2 - r1 + 1) > 200000:
            return f"'{sheet}'!{area}: values differ (range too large to itemize)"
        bad, total = [], 0
        for row in range(r1, r2 + 1):
            for col in range(c1, c2 + 1):
                total += 1
                want, got = gt.cell(row, col).value, proc.cell(row, col).value
                if not self.module.compare_cell_value(want, got):
                    bad.append((gt.cell(row, col).coordinate, want, got))
        shown = ", ".join(f"{c} expected {_short(w)} got {_short(g)}" if reveal else f"{c} got {_short(g)}"
                          for c, w, g in bad[:limit])
        return f"'{sheet}'!{area}: {len(bad)}/{total} cells wrong ({shown})"

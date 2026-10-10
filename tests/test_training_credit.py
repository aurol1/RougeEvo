import json
from pathlib import Path
import random
from tests.workspace import TemporaryDirectory
import unittest

import openpyxl

from rouge.core import Budget, Evaluation, TaskResult
from rouge.data import Case, OfficialScorer, Task
from rouge.evolution import bounded_training_trace, focused_evidence, trace_excerpt
from rouge.execution import SpreadsheetEvaluator
from rouge.response_search import training_quality

HERE = Path(__file__).resolve().parents[1]


@unittest.skipUnless((HERE / 'data/official_evaluator.py').is_file(),
                     'Prepare the pinned official evaluator to run scorer integration tests')
class TrainingCreditScorerTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.scorer = OfficialScorer(HERE / 'data/official_evaluator.py')
        self.answer = self.root / 'answer.xlsx'
        self.output = self.root / 'output.xlsx'
        self.case = Case(self.root / 'input.xlsx', self.answer)
        self.save(self.answer, [11, 22, 33, 44], other=True)
        self.save(self.output, [11, 999, 33, 999])

    def save(self, path, values, other=False):
        book = openpyxl.Workbook()
        book.active.title = 'S'
        for row, value in enumerate(values, 1):
            book['S'].cell(row, 1, value)
        if other:
            book.create_sheet('Other')['B1'] = 'PRIVATE_EXPECTED_VALUE'
        book.save(path)
        book.close()

    def task(self, targets=(('S', 'A1:A3'), ('S', 'A2:A4'))):
        return Task('train-id', 'Apply the requested transformation.', 'Cell-Level Manipulation',
                    "'S'!A1:A4", (self.case,), targets)

    def test_overlapping_target_ranges_count_union_and_keep_hard_failure(self):
        task = self.task()
        progress = self.scorer.training_progress(task, self.case, self.output)
        self.assertEqual(progress, {'available': True, 'matched': 2, 'total': 4, 'fraction': 0.5})
        self.assertFalse(self.scorer.compare(task, self.case, self.output))
        self.assertNotIn('999', json.dumps(progress))
        self.assertNotIn('expected', json.dumps(progress))
        self.assertNotIn('PRIVATE_EXPECTED_VALUE', json.dumps(progress))

    def test_passing_range_and_missing_sheet_are_both_counted(self):
        task = self.task((('S', 'A1'), ('S', 'A2:A4'), ('Other', 'B1')))
        progress = self.scorer.training_progress(task, self.case, self.output)
        self.assertEqual((progress['matched'], progress['total']), (2, 5))
        self.assertAlmostEqual(progress['fraction'], 0.4)

    def test_missing_or_corrupt_output_gets_zero_credit_not_unavailable(self):
        self.output.write_bytes(b'not a workbook')
        for path in (self.output, self.root / 'missing.xlsx'):
            with self.subTest(path=path.name):
                progress = self.scorer.training_progress(self.task(), self.case, path)
                self.assertTrue(progress['available'])
                self.assertEqual((progress['matched'], progress['total'], progress['fraction']), (0, 4, 0))

    def test_legacy_positions_and_whole_column_bound_use_same_union(self):
        task = self.task(None)
        progress = self.scorer.training_progress(task, self.case, self.output)
        self.assertEqual((progress['matched'], progress['total']), (2, 4))
        task = self.task((('S', 'A:A'), ('S', 'A1:A4')))
        progress = self.scorer.training_progress(task, self.case, self.output)
        self.assertEqual((progress['matched'], progress['total']), (2, 4))

    def test_official_numeric_equality_is_used_in_dense_credit(self):
        self.save(self.output, [11.0, 22.0, 33.0, 44.0])
        progress = self.scorer.training_progress(self.task(), self.case, self.output)
        self.assertEqual(progress['fraction'], 1)
        self.assertTrue(self.scorer.compare(self.task(), self.case, self.output))


class FakeRunner:
    def recalculate(self, workspace, budget):
        pass


class AggregateScorer:
    def __init__(self):
        self.credit_calls = []

    def check(self, task, case, path, **kwargs):
        return False, 'wrong cells'

    def training_progress(self, task, case, path):
        self.credit_calls.append((task.id, case))
        return {'available': True, 'matched': 1, 'total': 2 if case.answer.name == 'a.xlsx' else 4,
                'fraction': 0.5 if case.answer.name == 'a.xlsx' else 0.25}


class TrainingCreditBoundaryTests(unittest.TestCase):
    def evaluator(self, enabled=True):
        evaluator = SpreadsheetEvaluator.__new__(SpreadsheetEvaluator)
        evaluator.runner = FakeRunner()
        evaluator.budget = Budget(1000, tool_price=10)
        evaluator.explain_ids = {'train-id'}
        evaluator.config = {'training_cell_credit': enabled, 'reveal_expected': False}
        evaluator.scorer = AggregateScorer()
        return evaluator

    def task(self, task_id='train-id'):
        cases = (Case(Path('in1.xlsx'), Path('a.xlsx')), Case(Path('in2.xlsx'), Path('b.xlsx')))
        return Task(task_id, 'A task.', 'Cell-Level Manipulation', "'S'!A1", cases)

    def test_dense_score_is_macro_case_credit_and_does_not_change_verdict(self):
        evaluator, sink = self.evaluator(), {}
        cases, _ = evaluator._score(self.task(), Path('.'), sink)
        self.assertEqual(cases, [False, False])
        self.assertEqual(sink['score'], 0.375)
        self.assertEqual(len(evaluator.scorer.credit_calls), 2)
        self.assertEqual(evaluator.budget.spent, 10)

    def test_validation_and_disabled_credit_do_not_read_dense_answers(self):
        for enabled, task_id in ((True, 'validation-id'), (False, 'train-id')):
            with self.subTest(enabled=enabled, task_id=task_id):
                evaluator, sink = self.evaluator(enabled), {}
                evaluator._score(self.task(task_id), Path('.'), sink)
                self.assertEqual(sink, {})
                self.assertEqual(evaluator.scorer.credit_calls, [])

    def test_training_quality_has_hard_fallback_without_partial_credit(self):
        result = Evaluation('build', [
            TaskResult('a', False, [False], '', '', training_progress={'available': True, 'score': 0.5}),
            TaskResult('b', True, [True], '', '', training_progress={'available': False}),
            TaskResult('c', False, [False], '', '')])
        self.assertEqual(result.score, 1 / 3)
        self.assertEqual(training_quality(result), 0.5)

    def test_training_trace_preserves_code_errors_but_excludes_observations(self):
        trace = [
            {'step': -1, 'stage': 'input_observation', 'facts': 'INPUT_FACTS_SHOULD_NOT_ENTER_CODE_TRACE'},
            {'step': 0, 'response': '```python\nprint("CODE_FIRST")\n```', 'returncode': 0, 'feedback': 'ok'},
            {'step': 1, 'response': '```python\nwrong_call()\n```', 'returncode': 1, 'feedback': 'Traceback ERROR_MARKER'},
            {'step': 2, 'response': 'DONE', 'returncode': 0, 'feedback': 'ok'}]
        encoded = bounded_training_trace(trace, 1600)
        self.assertLessEqual(len(encoded), 1600)
        self.assertEqual(len(json.loads(encoded)), 3)
        excerpt = trace_excerpt(encoded, 1200)
        self.assertIn('CODE_FIRST', excerpt)
        self.assertIn('ERROR_MARKER', excerpt)
        self.assertNotIn('INPUT_FACTS_SHOULD_NOT_ENTER_CODE_TRACE', encoded)
        self.assertNotEqual(excerpt, 'DONE')

    def test_escaped_trace_is_bounded_valid_json_and_focused_evidence_has_no_task_id(self):
        trace = [{'step': 0, 'response': '"\\\n' * 10000, 'feedback': '\n"' * 10000, 'returncode': 1}]
        encoded = bounded_training_trace(trace, 512)
        self.assertLessEqual(len(encoded), 512)
        self.assertIsInstance(json.loads(encoded), list)
        task = self.task('PRIVATE_TASK_ID')
        evaluation = Evaluation('build', [TaskResult(task.id, False, [False], 'failed',
            json.dumps([{'step': 0, 'response': '```python\nprint("TRACE_CODE")\n```',
                         'returncode': 1, 'feedback': 'Traceback'}]))])
        evidence = focused_evidence({task.id: task}, evaluation, random.Random(0), 1200)
        self.assertLessEqual(len(evidence), 1200)
        self.assertIn('TRACE_CODE', evidence)
        self.assertNotIn('PRIVATE_TASK_ID', evidence)


if __name__ == '__main__':
    unittest.main()
